#!/usr/bin/env python3
"""中間発表スライドの雛形を生成する。

内容の正は docs/presentation_midterm.md。この台本に沿って pptx を組み立てる。
生成物: docs/export/中間発表_柴尾.pptx

    python3 -m pip install --user python-pptx
    python3 scripts/build_midterm_pptx.py
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXPORT = os.path.join(ROOT, "docs", "export")
OUT = os.path.join(EXPORT, "中間発表_柴尾.pptx")

# 図・スクリーンショットの場所。無ければ枠だけ描いて先に進む。
DESKTOP = os.path.expanduser("~/Desktop/卒業研究")
IMG = {
    # scripts/render_figures.sh で mermaid から生成した高解像度版
    "asis": os.path.join(EXPORT, "fig_asis.png"),
    "tobe": os.path.join(EXPORT, "fig_tobe.png"),
    "bd": os.path.join(DESKTOP, "6:25", "ブロックダイアグラム.png"),
    "fp": os.path.join(DESKTOP, "6:25", "フロントパネル.png"),
}

FONT = "Yu Gothic"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x0E, 0x7C, 0x86)
RED = RGBColor(0xC0, 0x39, 0x2B)
GRAY = RGBColor(0x59, 0x5A, 0x5C)
LIGHT = RGBColor(0xEF, 0xF3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)


def set_font(run, size=18, bold=False, color=None, font=FONT):
    """欧文・日本語の両方に同じフォントを当てる。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.makeelement(qn(tag), {"typeface": font})
        rPr.append(el)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return box, tf


def para(tf, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT,
         space_after=6, level=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color)
    return p


def slide_number(slide, n):
    box, tf = textbox(slide, SW - Inches(1.0), SH - Inches(0.55), Inches(0.7), Inches(0.4))
    para(tf, str(n), size=14, color=GRAY, align=PP_ALIGN.RIGHT, first=True)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def base_slide(prs, title=None, number=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙
    if title is not None:
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        bar.shadow.inherit = False
        tf = bar.text_frame
        tf.margin_left = Inches(0.5)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, title, size=30, bold=True, color=WHITE, first=True)
    if number is not None:
        slide_number(s, number)
    return s


def bullets(slide, items, x=Inches(0.8), y=Inches(1.4), w=Inches(11.7), h=Inches(5.4),
            size=20, gap=12):
    """items: (テキスト, 階層, 強調) のリスト。文字列だけでも可。"""
    box, tf = textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        if isinstance(item, str):
            text, level, emph = item, 0, False
        else:
            text, level, emph = (list(item) + [0, False])[:3]
        color = ACCENT if emph else RGBColor(0x22, 0x22, 0x22)
        prefix = "" if level == 0 else "− "
        para(tf, prefix + text, size=size - 2 * level, bold=emph,
             color=color, space_after=gap, level=level, first=(i == 0))
    return box


def card(slide, x, y, w, h, heading, lines, accent=ACCENT, head_size=20, body_size=16):
    """見出し付きの囲みブロック。"""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = accent
    box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.18)
    para(tf, heading, size=head_size, bold=True, color=accent, first=True, space_after=8)
    for line in lines:
        para(tf, line, size=body_size, color=RGBColor(0x22, 0x22, 0x22), space_after=5)
    return box


def table(slide, rows, x, y, w, col_widths=None, row_h=Inches(0.45),
          size=15, head_size=15, highlight_row=None, head_h=None):
    n_rows, n_cols = len(rows), len(rows[0])
    head_h = head_h or Inches(0.45)
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, head_h + row_h * (n_rows - 1))
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for r, row in enumerate(rows):
        tbl.rows[r].height = head_h if r == 0 else row_h
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            is_head = r == 0
            emph = highlight_row is not None and r == highlight_row
            para(tf, str(val), size=head_size if is_head else size,
                 bold=is_head or emph,
                 color=WHITE if is_head else (RED if emph else RGBColor(0x22, 0x22, 0x22)),
                 first=True, space_after=0)
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = NAVY
            elif emph:
                cell.fill.fore_color.rgb = RGBColor(0xFD, 0xEC, 0xEA)
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
    return shape


def picture(slide, key, x, y, max_w, max_h, caption=None):
    """縦横比を保ったまま max_w × max_h の枠に収め、枠の中央に置く。"""
    path = IMG.get(key)
    if path and os.path.exists(path):
        pic = slide.shapes.add_picture(path, x, y)
        scale = min(max_w / pic.width, max_h / pic.height)
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
        pic.left = int(x + (max_w - pic.width) / 2)
        pic.top = int(y + (max_h - pic.height) / 2)
        if caption:
            _, tf = textbox(slide, x, pic.top + pic.height + Inches(0.08),
                            max_w, Inches(0.35))
            para(tf, caption, size=13, color=GRAY, align=PP_ALIGN.CENTER, first=True)
        return pic
    else:
        ph = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, max_w, max_h)
        ph.fill.solid()
        ph.fill.fore_color.rgb = LIGHT
        ph.line.color.rgb = GRAY
        ph.shadow.inherit = False
        tf = ph.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, f"[画像を挿入: {os.path.basename(path or key)}]",
             size=16, color=GRAY, align=PP_ALIGN.CENTER, first=True)
    if caption:
        _, tf = textbox(slide, x, y + max_h + Inches(0.05), max_w, Inches(0.35))
        para(tf, caption, size=13, color=GRAY, align=PP_ALIGN.CENTER, first=True)


def roadmap(slide, x, y, w, h):
    """発展ロードマップ。Mermaidのtimelineは日本語が崩れるため図形で描く。"""
    stages = [
        ("現在", ["LabVIEWとLLMのJSON接続基盤", "サーバー側の会話制御", "（ここまで実装済み）"], True),
        ("次段階", ["司会AIの導入", "ペルソナの動的な切り替え", "（本研究で取り組む範囲）"], False),
        ("中期", ["研究室の既存研究とLLMを", "繋ぐゲートウェイ化", "センサーデータとの統合"], False),
        ("長期［ビジョン］", ["客観データと主観データを", "組み合わせた評価空間"], False),
    ]
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y + h - Inches(0.42),
                                   w, Inches(0.32))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(0xC8, 0xD3, 0xDF)
    arrow.line.fill.background()
    arrow.shadow.inherit = False

    gap = Inches(0.2)
    bw = int((w - gap * (len(stages) - 1)) / len(stages))
    for i, (head, lines, done) in enumerate(stages):
        bx = x + i * (bw + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, bw,
                                     h - Inches(0.55))
        box.fill.solid()
        box.fill.fore_color.rgb = ACCENT if done else LIGHT
        box.line.color.rgb = ACCENT if done else GRAY
        box.line.width = Pt(2 if done else 1)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.12)
        para(tf, head, size=19, bold=True, color=WHITE if done else NAVY,
             align=PP_ALIGN.CENTER, first=True, space_after=6)
        for ln in lines:
            para(tf, ln, size=13, color=WHITE if done else RGBColor(0x22, 0x22, 0x22),
                 align=PP_ALIGN.CENTER, space_after=3)


def video_placeholder(slide, x, y, w, h, label, howto):
    """動画を後から差し込むための枠。撮影後にここへ挿入する。"""
    ph = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    ph.fill.solid()
    ph.fill.fore_color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    ph.line.color.rgb = RED
    ph.line.width = Pt(2.5)
    ph.line.dash_style = 4  # 破線
    ph.shadow.inherit = False
    tf = ph.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "▶  " + label, size=26, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, first=True, space_after=10)
    para(tf, howto, size=14, color=RGBColor(0xDD, 0xDD, 0xDD), align=PP_ALIGN.CENTER)
    para(tf, "挿入 → ビデオ → このデバイス（埋め込み） / 再生タブ → 開始「自動」",
         size=12, color=RGBColor(0xAA, 0xAA, 0xAA), align=PP_ALIGN.CENTER)
    return ph


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    n = 0

    # ---------- 1. 表紙 ----------
    s = base_slide(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(4.3))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    _, tf = textbox(s, Inches(1.0), Inches(1.2), Inches(11.3), Inches(2.6))
    para(tf, "LabVIEWと大規模言語モデルを接続する", size=38, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, first=True, space_after=10)
    para(tf, "マルチエージェント対話基盤の構築", size=38, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)
    _, tf = textbox(s, Inches(1.0), Inches(4.8), Inches(11.3), Inches(2.0))
    para(tf, "情報システム工学科　松田研究室", size=22, color=NAVY,
         align=PP_ALIGN.CENTER, first=True, space_after=8)
    para(tf, "（学籍番号）　柴尾 海渡", size=26, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
         space_after=14)
    para(tf, "2026年　月　日　　卒業研究 中間発表", size=18, color=GRAY, align=PP_ALIGN.CENTER)
    notes(s, "松田研究室の柴尾です。「LabVIEWと大規模言語モデルを接続する、"
             "マルチエージェント対話基盤の構築」について発表します。")

    # ---------- 2. 目次 ----------
    n += 1
    s = base_slide(prs, "目次", n)
    bullets(s, [
        "1.  研究背景",
        "2.  問題点",
        "3.  研究目的",
        "4.  システム構成",
        "5.  現在までの進捗　― 動作映像 ―",
        "6.  定量評価",
        "7.  今後の方針と本研究の位置づけ",
    ], y=Inches(1.7), size=26, gap=20)
    notes(s, "こちらの流れで発表します。中盤で実際に動いている様子を映像でお見せします。")

    # ---------- 3. 研究背景 ----------
    n += 1
    s = base_slide(prs, "研究背景", n)
    bullets(s, [
        "リハビリテーション評価は、療法士が目視で採点する方式が主流",
        "高齢化により対象者は増加し、評価の負担も増している",
        "センサーにより客観データは取得できるようになった",
        ("指の屈曲角 ・ 運動回数 ・ 脈波 など", 1),
    ], y=Inches(1.3), h=Inches(2.2), size=20)
    card(s, Inches(0.8), Inches(3.5), Inches(5.5), Inches(1.5), "客観データ（測れる）",
         ["センサーで数値として取得できる", "屈曲角・運動回数・脈波"], accent=ACCENT)
    card(s, Inches(7.0), Inches(3.5), Inches(5.5), Inches(1.5), "主観（測れていない）",
         ["「動きやすくなった実感はあるか」", "「続けたいと思えるか」"], accent=RED)
    _, tf = textbox(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.0))
    para(tf, "両方が揃って、はじめて評価になる", size=24, bold=True, color=NAVY,
         align=PP_ALIGN.CENTER, first=True)
    notes(s,
          "現在のリハビリテーション評価は、療法士が目視で採点する方式が主流です。"
          "高齢化により対象者は増え続けており、評価の負担も増しています。\n"
          "一方、指の曲がり具合や運動回数、脈波といったセンサーデータは取得できるようになってきました。"
          "ただ、それだけでは足りません。「動くようになった実感があるか」「続けたいと思えるか」"
          "といった本人の主観は、センサーでは測れないからです。\n"
          "この主観を自然な会話から引き出し、客観データと組み合わせて評価したい、"
          "というのが本研究の出発点です。")

    # ---------- 4. 問題点 ----------
    n += 1
    s = base_slide(prs, "問題点", n)
    card(s, Inches(0.7), Inches(1.35), Inches(5.9), Inches(2.5),
         "問題点 ①　主観を引き出す仕組みがない",
         ["・質問紙は形式が固定されており、",
          "　その場の話に応じた掘り下げができない",
          "・人が聞き取ると、聞き手の関心によって",
          "　内容が偏ってしまう"], accent=RED)
    card(s, Inches(6.9), Inches(1.35), Inches(5.7), Inches(2.5),
         "問題点 ②　画面変更のたびに作り直し",
         ["・計測画面はLabVIEWで作られており、",
          "　評価内容が変わるたびに画面も変わる",
          "・別に専用のWeb画面を作ると、",
          "　変更のたびに二重の修正が必要になる"], accent=RED)
    _, tf = textbox(s, Inches(0.7), Inches(4.3), Inches(11.9), Inches(2.0))
    para(tf, "解くべき課題", size=20, bold=True, color=GRAY, first=True, space_after=12)
    para(tf, "偏りなく主観を引き出せて、かつ画面の変更に強い仕組みが要る",
         size=26, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    notes(s,
          "問題点を2つに整理しました。\n"
          "1つ目は、主観を引き出す仕組みがないことです。質問紙は形式が固定されているため、"
          "その場の話に応じて掘り下げることができません。かといって人が聞き取ると、"
          "聞き手の関心によって内容が偏ってしまいます。\n"
          "2つ目は開発面の問題です。この研究室では計測系の画面をLabVIEWで作っていますが、"
          "評価内容が変われば画面も変わります。別に専用のWeb画面を作ってしまうと、"
          "変更のたびに二重の修正が必要になります。")

    # ---------- 5. 研究目的 ----------
    n += 1
    s = base_slide(prs, "研究目的", n)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3),
                             Inches(11.7), Inches(1.5))
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "LabVIEWから大規模言語モデル（LLM）を利用できる接続基盤を構築し、",
         size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, space_after=6)
    para(tf, "複数のAIが対話する形で主観情報を引き出せるようにする",
         size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for i, (head, body) in enumerate([
        ("① 画面変更に強い構成", ["LabVIEWをクライアントとし、", "画面の変更をLabVIEW側だけで", "完結できるようにする"]),
        ("② 複数AIの会話制御", ["複数のAIが順番に発言する", "会話をサーバー側で制御する"]),
        ("③ 引き継げる環境", ["研究室の他のPCでも", "同じように動く構成にする"]),
    ]):
        card(s, Inches(0.8) + i * Inches(4.0), Inches(3.3), Inches(3.7), Inches(2.4),
             head, body, head_size=19)
    notes(s,
          "研究目的です。LabVIEWから大規模言語モデルを利用できる接続基盤を構築し、"
          "複数のAIが対話する形で主観情報を引き出せるようにすることを目指します。\n"
          "具体的には3点です。1点目、LabVIEWをクライアントにして、画面の変更をLabVIEW側だけで"
          "完結できるようにすること。2点目、複数のAIが順番に発言する会話をサーバー側で制御すること。"
          "3点目、研究室の他のPCでも動く、引き継ぎ可能な構成にすることです。")

    # ---------- 6. システム構成 ----------
    n += 1
    s = base_slide(prs, "システム構成", n)
    # As-Is図は横長（約6.8:1）なので全幅に置く
    picture(s, "asis", Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.9),
            caption="現在のシステム構成（As-Is）")
    for i, (head, body) in enumerate([
        ("LabVIEW ⇄ サーバー", ["HTTP / JSON", "文字列をJSONにして送るだけ"]),
        ("サーバー ⇄ LLM", ["HTTPS", "APIキーはサーバー内のみ保持し", "LabVIEWには渡さない"]),
        ("実行環境", ["Docker", "どのPCでも同じ環境で起動できる"]),
    ]):
        card(s, Inches(0.7) + i * Inches(4.1), Inches(4.0), Inches(3.8), Inches(1.8),
             head, body, head_size=19, body_size=15)
    notes(s,
          "LabVIEWをクライアント、サーバーをLLMとの仲介役とする構成にしました。\n"
          "LabVIEWからは文字列をJSON形式にしてHTTPで送るだけです。サーバーがAPIキーを付けて"
          "LLMに問い合わせ、返ってきた回答を整形してLabVIEWに返します。\n"
          "ポイントは、APIキーをLabVIEW側に一切持たせていないことです。"
          "またサーバーはDocker上で動くようにしてあり、どのPCでも同じ環境で起動できます。")

    # ---------- 7. 設計方針 ----------
    n += 1
    s = base_slide(prs, "設計方針 ― LabVIEW側の実装を増やさない", n)
    _, tf = textbox(s, Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.5))
    para(tf, "サーバー側で吸収することで、LabVIEW側は最小限の部品で済むように設計した",
         size=18, color=GRAY, first=True)
    table(s, [
        ["工夫", "内容", "LabVIEW側での効果"],
        ["① 平坦なJSON", "入れ子構造をやめ、フラットな形にする",
         "クラスタを1つ作るだけで送受信できる"],
        ["② 常にHTTP 200", "成否は本文の ok フラグで表す",
         "ステータス分岐が不要。Case Structure 1つで済む"],
        ["③ 会話を1本の文字列でも返す", "発言の配列に加え、連結済みの文字列も返す",
         "配列を解析せずに会話全体を表示できる"],
    ], Inches(0.8), Inches(1.9), Inches(11.7), col_widths=[3, 5, 5.5],
        row_h=Inches(0.95), head_h=Inches(0.5), size=16, head_size=17)
    _, tf = textbox(s, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.0))
    para(tf, "→ 複雑さをサーバー側に寄せることで、画面の変更にLabVIEW側だけで追従できる",
         size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True)
    notes(s,
          "設計で一番気を使ったのは、LabVIEW側の実装をいかに増やさないかです。3つの工夫をしました。\n"
          "1つ目、JSONを入れ子にせず平坦な形にしました。LabVIEWはクラスタのラベル名がそのまま"
          "JSONのキーになるので、平坦なほど扱いが簡単になります。\n"
          "2つ目、エラーのときもHTTPステータスは200を返し、成否は本文のokというフラグで表すようにしました。"
          "これでLabVIEW側はステータスコードの分岐をせず、Case Structure 1つで済みます。\n"
          "3つ目、複数AIの会話は、発言の配列だけでなく、それらを1本につないだ文字列も同時に返すようにしました。")

    # ---------- 8. マルチエージェント化の段階 ----------
    n += 1
    s = base_slide(prs, "マルチエージェント化の段階", n)
    picture(s, "tobe", Inches(0.7), Inches(1.2), Inches(6.6), Inches(4.6),
            caption="将来構成（To-Be）")
    table(s, [
        ["段階", "内容", "状態"],
        ["A", "LabVIEW主導でAIを1体ずつ呼ぶ", "実現可能"],
        ["B", "サーバーが発言順を制御しまとめて返す", "実装完了 ←現在地"],
        ["C", "司会AIが発言順を動的に決める", "次の課題"],
    ], Inches(7.7), Inches(1.5), Inches(4.9), col_widths=[1, 5, 2.4],
        row_h=Inches(0.8), head_h=Inches(0.45), size=14, head_size=15, highlight_row=2)
    card(s, Inches(7.7), Inches(4.3), Inches(4.9), Inches(1.5), "段階Cへの備え",
         ["発言順を「決める処理」と「実行する処理」を",
          "分けて実装した。司会AIの導入時は、",
          "決める部分だけを差し替えれば済む。"], body_size=14)
    notes(s,
          "マルチエージェント化は3段階で進めています。\n"
          "段階Aは、LabVIEW側が1体ずつAIを呼ぶ方式です。今の仕組みでも可能ですが、"
          "会話が増えるほどLabVIEW側が複雑になります。\n"
          "そこで段階Bとして、発言順の制御をサーバー側に持たせました。LabVIEWは1回リクエストを"
          "送るだけで、複数のAIの発言がまとめて返ってきます。ここまでが実装完了しています。\n"
          "段階Cが今後の課題で、司会役のAIが文脈に応じて次の発話者を決める形にします。"
          "今回の実装では、この発言順を決める部分だけを差し替えれば済む構造にしてあります。")

    # ---------- 9. 動作映像① ----------
    n += 1
    s = base_slide(prs, "動作映像 ①　単体AIとの往復", n)
    video_placeholder(s, Inches(1.6), Inches(1.35), Inches(10.1), Inches(4.6),
                      "動画① をここに挿入（30秒）",
                      "撮影手順: docs/presentation_video_guide.md の 3-3")
    _, tf = textbox(s, Inches(1.6), Inches(6.15), Inches(10.1), Inches(0.6))
    para(tf, "LabVIEWで入力 → サーバー経由でLLMへ → 返信がLabVIEWに表示される",
         size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True)
    notes(s,
          "ここから実際に動いている様子をお見せします。まずは基本となる、1体のAIとの往復です。\n"
          "（再生）LabVIEWのフロントパネルに文章を入力して実行します。右側がサーバーのログで、"
          "リクエストが届いているのが見えます。数秒後、AIの返信がLabVIEW側に表示されます。\n"
          "ここまでが基本形で、この上に複数AIの仕組みを載せています。")

    # ---------- 10. 動作映像② ----------
    n += 1
    s = base_slide(prs, "動作映像 ②　複数AIの会話", n)
    video_placeholder(s, Inches(0.6), Inches(1.25), Inches(7.2), Inches(4.3),
                      "動画② をここに挿入（60秒）",
                      "撮影手順: docs/presentation_video_guide.md の 3-4")
    _, tf = textbox(s, Inches(0.6), Inches(5.7), Inches(7.2), Inches(1.3))
    para(tf, "入力（LabVIEWから1回送信するだけ）", size=15, bold=True, color=GRAY,
         first=True, space_after=6)
    para(tf, "「リハビリのトレーニングを終えました。指が少し動きやすく"
             "なった気がして、続けられそうです。」", size=16, color=RGBColor(0x22, 0x22, 0x22))
    table(s, [
        ["順", "AI", "発言の要旨"],
        ["1", "Alpha", "素晴らしいですね。改善を実感できることは励みになります"],
        ["2", "Beta", "良い兆候です。ただし無理をせず、専門家の指導を仰ぐことも大切です"],
        ["3", "Gamma", "その体験を他の挑戦にも生かすと、日常生活にも可能性が広がります"],
    ], Inches(8.1), Inches(1.45), Inches(4.6), col_widths=[0.7, 1.5, 6],
        row_h=Inches(0.95), head_h=Inches(0.42), size=13, head_size=14, highlight_row=2)
    card(s, Inches(8.1), Inches(5.0), Inches(4.6), Inches(1.9), "注目点", [
        "2体目が「ただし」と受けている。",
        "同じ質問への答えを並べたのではなく、",
        "前の発言を読んだうえで発言しているため、",
        "内容が重複せず視点が加わっている。"], accent=RED, head_size=18, body_size=14)
    notes(s,
          "こちらが今回の中心です。LabVIEWから1回送信するだけで、3体のAIが順番に発言します。\n"
          "（再生）入力するのは、患者さんの感想を想定した1文だけです。実行すると、サーバー側で"
          "AIが順番に呼ばれていきます。ログを見ると、3回のやり取りが順に行われているのが分かります。\n"
          "結果がこちらです。注目していただきたいのは、2体目が「ただし」と受けている点です。"
          "3体に同じ質問を投げて答えを並べたのではなく、前の発言を読んだうえで発言しているので、"
          "内容が重複せず、慎重な視点や発想を広げる視点が加わっています。\n"
          "これは各AIに直前までの発言を文脈として渡し、既出の意見を繰り返さないよう指示しているためです。"
          "将来ここに専門領域ごとの役割を割り当てることで、多角的な評価につなげられると考えています。")

    # ---------- 11. 動作映像③ ----------
    n += 1
    s = base_slide(prs, "動作映像 ③　別PCでの起動（移植性の検証）", n)
    video_placeholder(s, Inches(0.7), Inches(1.3), Inches(8.0), Inches(4.7),
                      "動画③ をここに挿入（40秒）",
                      "撮影手順: docs/presentation_video_guide.md の 3-5")
    card(s, Inches(9.1), Inches(1.4), Inches(3.5), Inches(4.5), "検証内容", [
        "研究室の資産として",
        "引き継げることが要件。",
        "",
        "Dockerで環境ごと固めた。",
        "",
        "・Node.js 不要",
        "・Git 不要",
        "・ZIP展開 → コマンド1行",
        "",
        "他メンバーのPCでも、",
        "数年後の引き継ぎでも、",
        "同じ環境を再現できる。"], body_size=14)
    notes(s,
          "3つ目は、このサーバーが他のPCでも動くかという検証です。\n"
          "研究室の資産として引き継げることが重要なので、Dockerで環境ごと固めました。\n"
          "（再生）リポジトリをZIPでダウンロードした状態から、コマンド1行で起動しています。"
          "Node.jsもGitもインストールしていません。起動後、そのままAIの応答まで通ることが確認できます。\n"
          "これにより、他のメンバーのPCでも、また数年後に引き継ぐ場合でも、同じ環境を再現できます。")

    # ---------- 12. 定量評価 ----------
    n += 1
    s = base_slide(prs, "定量評価 ― 応答時間", n)
    table(s, [
        ["条件", "応答時間", "内訳"],
        ["AI 1体", "約 1.2 〜 2.6 秒", "サーバー処理 数ms ＋ LLM応答待ち"],
        ["AI 2体", "約 1.5 〜 4.7 秒", "サーバー処理 数ms ＋ LLM応答待ち"],
        ["AI 3体", "約 6.3 秒", "サーバー処理 7ms ＋ LLM応答待ち 6.3秒"],
    ], Inches(0.8), Inches(1.3), Inches(11.7), col_widths=[2.5, 3.5, 6],
        row_h=Inches(0.75), head_h=Inches(0.5), size=17, head_size=17, highlight_row=3)
    card(s, Inches(0.8), Inches(4.3), Inches(3.7), Inches(1.8), "自作部分の負荷は無視できる",
         ["遅延のほぼ全てがLLM側の応答待ち。", "自作サーバーの処理は7ミリ秒。"],
         accent=ACCENT, head_size=17, body_size=14)
    card(s, Inches(4.8), Inches(4.3), Inches(3.7), Inches(1.8), "実用上は3体程度",
         ["AIの数にほぼ比例して伸びる。", "会話として成立する範囲で", "3体程度が現実的。"],
         accent=ACCENT, head_size=17, body_size=14)
    card(s, Inches(8.8), Inches(4.3), Inches(3.7), Inches(1.8), "タイムアウト対策",
         ["LabVIEWの通信タイムアウトは", "既定10秒のため3体では危険。", "延長して対応済み。"],
         accent=RED, head_size=17, body_size=14)
    notes(s,
          "定量評価として、応答時間を計測しました。\n"
          "AIが3体の場合で約6.3秒です。内訳を見ると、自作サーバー側の処理は7ミリ秒で、"
          "残りはすべてLLM側の応答待ちでした。つまり遅延の原因はほぼ外部APIであり、"
          "自作部分のオーバーヘッドは無視できる水準です。\n"
          "AIの数にほぼ比例して伸びるため、実際の対話では3体程度が現実的だと考えています。\n"
          "なお、LabVIEWの通信タイムアウトは既定で10秒のため、3体だと危険な水準です。"
          "これは実装側で延長して対応しています。")

    # ---------- 13. 今後の方針 ----------
    n += 1
    s = base_slide(prs, "今後の方針", n)
    roadmap(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(2.2))
    _, tf = textbox(s, Inches(0.7), Inches(3.75), Inches(11.9), Inches(0.5))
    para(tf, "本研究で取り組む範囲（次段階）", size=20, bold=True, color=NAVY, first=True)
    items = [
        ("段階C：司会AIの導入", ["文脈に応じて次の発話者を決め、", "発言の偏りを制御する"]),
        ("役割（ペルソナ）の設計", ["どのような視点のAIを揃えるかは", "臨床的な知見が必要"]),
        ("通信の暗号化", ["実データを扱う段階では", "HTTPS化が必須"]),
    ]
    for i, (head, body) in enumerate(items):
        card(s, Inches(0.7) + i * Inches(4.1), Inches(4.35), Inches(3.8), Inches(1.6),
             head, body, head_size=17, body_size=14)
    _, tf = textbox(s, Inches(0.7), Inches(6.15), Inches(11.9), Inches(0.5))
    para(tf, "中期以降は本研究の完了範囲の外にあり、発展の方向性として示している",
         size=15, color=GRAY, first=True)
    notes(s,
          "今後の方針です。\n"
          "本研究で取り組む範囲は、図の「次段階」までです。まず段階Cとして司会役のAIを導入し、"
          "文脈に応じて次に話すAIを決めることで、発言の偏りを制御したいと考えています。\n"
          "ただし、どのような視点のAIを揃えるべきか、どう質問を投げれば主観を引き出せるかは、"
          "臨床的な知見が必要な部分です。ここは相談しながら進めたいと考えています。\n"
          "中期以降は本研究の完了範囲の外で、発展の方向性としてお示ししているものです。"
          "その中期の位置づけを、次のスライドで説明します。")

    # ---------- 14. 本研究の位置づけ ----------
    n += 1
    s = base_slide(prs, "本研究の位置づけ ― 既存研究とLLMを繋ぐゲートウェイ", n)
    _, tf = textbox(s, Inches(0.7), Inches(1.05), Inches(11.9), Inches(0.5))
    para(tf, "研究室の既存研究はすでにLabVIEWで動いている。"
             "本研究はその共通の出口としてLLMを繋ぐ役割を担える。",
         size=17, color=GRAY, first=True)

    cols = [
        (Inches(0.6), Inches(3.5), "松田研の既存研究", [
            "・リハビリ用デバイス", "　（力覚・屈曲データ）",
            "・ハードウェア制御", "　（PWM / AD・DA変換）",
            "・電動車椅子・計測系",
            "", "いずれもLabVIEWが窓口"], GRAY),
        (Inches(4.9), Inches(3.6), "本研究：接続基盤", [
            "・LabVIEWから1回投げるだけ",
            "・サーバーが会話を制御",
            "・LabVIEW側に追加実装は不要",
            "", "計測系を作り変えずに", "LLMを利用できる"], ACCENT),
        (Inches(9.3), Inches(3.4), "LLM / マルチエージェント", [
            "・複数AIによる対話",
            "・主観情報の引き出し",
            "・司会AIによる発言順の制御",
            "", "（外部API / 将来は院内）"], NAVY),
    ]
    for cx, cw, head, lines, accent in cols:
        card(s, cx, Inches(1.65), cw, Inches(2.6), head, lines,
             accent=accent, head_size=18, body_size=14)
    for ax in (Inches(4.25), Inches(8.65)):
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, Inches(2.7), Inches(0.55),
                                Inches(0.45))
        ar.fill.solid()
        ar.fill.fore_color.rgb = ACCENT
        ar.line.fill.background()
        ar.shadow.inherit = False

    band = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.55),
                              Inches(12.1), Inches(1.1))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "本研究は、研究室の既存研究にLLMを接続する共通の入口（ゲートウェイ）になり得る",
         size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)

    _, tf = textbox(s, Inches(0.6), Inches(5.85), Inches(12.1), Inches(0.6))
    para(tf, "※ センサーデータとの統合は中期の目標。本研究の完了範囲は接続基盤の構築まで。",
         size=15, color=RED, align=PP_ALIGN.CENTER, first=True)
    notes(s,
          "最後に、本研究の位置づけを整理します。\n"
          "この研究室の既存研究は、リハビリ用デバイスにしてもハードウェア制御にしても、"
          "いずれもLabVIEWが窓口になっています。\n"
          "本研究が作ったのは、そのLabVIEWから1回投げるだけでLLMを使える接続基盤です。"
          "つまり、既存の計測系を作り変えずにLLMを利用できるようになります。\n"
          "この意味で本研究は、研究室の既存研究にLLMを接続する共通の入口、"
          "ゲートウェイになり得ると考えています。\n"
          "なお、センサーデータとの統合そのものは中期の目標であり、"
          "本研究の完了範囲は接続基盤の構築までです。\n"
          "以上で発表を終わります。ありがとうございました。")

    # ---------- 15. 結び ----------
    s = base_slide(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.7), SW, Inches(2.1))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "ご清聴ありがとうございました", size=40, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, first=True)

    # ---------- 補足スライド ----------
    s = base_slide(prs, "補足資料")
    _, tf = textbox(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5))
    for i, t in enumerate([
        "補1．LabVIEW ブロックダイアグラム",
        "補2．JSON の具体例",
        "補3．自作部分と OSS の線引き",
        "補4．発言順の制御",
        "補5．堅牢性への配慮",
        "補6．開発体制",
    ]):
        para(tf, t, size=22, color=NAVY, first=(i == 0), space_after=14)

    # 補1
    s = base_slide(prs, "補1．LabVIEW ブロックダイアグラム")
    picture(s, "bd", Inches(1.2), Inches(1.25), Inches(10.9), Inches(4.9),
            caption="LabVIEW標準の部品のみで構成（HTTP Client VIs / JSON VIs）。追加ツールキットは不使用")

    # 補2
    s = base_slide(prs, "補2．JSON の具体例")
    card(s, Inches(0.7), Inches(1.3), Inches(5.9), Inches(2.4), "リクエスト（LabVIEW → サーバー）", [
        '{', '  "text": "リハビリを終えました…",', '  "agentIds": "alpha,beta,gamma",',
        '  "rounds": 1', '}'], head_size=17, body_size=14)
    card(s, Inches(6.9), Inches(1.3), Inches(5.7), Inches(2.4), "レスポンス（サーバー → LabVIEW）", [
        '{', '  "ok": true,', '  "transcript": "Alpha: …\\nBeta: …",',
        '  "error": ""', '}'], head_size=17, body_size=14)
    _, tf = textbox(s, Inches(0.7), Inches(4.0), Inches(11.9), Inches(2.4))
    para(tf, "設計上の要点", size=20, bold=True, color=NAVY, first=True, space_after=10)
    for t in [
        "・入れ子にしない。LabVIEWのクラスタのラベル名とキー名が1対1で対応する",
        "・エラー時もHTTPは200を返す。成否は ok で判定するためCase Structureが1つで済む",
        "・発言の配列も返しているが、連結済みの transcript だけでも会話全体を表示できる",
    ]:
        para(tf, t, size=17, space_after=8)

    # 補3
    s = base_slide(prs, "補3．自作部分と OSS の線引き")
    table(s, [
        ["レイヤ", "使用技術", "区分"],
        ["LabVIEW側 VI", "LabVIEW標準（HTTP Client VIs / JSON VIs）", "自作VI"],
        ["実行環境", "Docker", "OSS"],
        ["Webフレームワーク・言語", "Next.js / TypeScript / Node.js", "OSS"],
        ["LLM接続", "各社公式SDK", "OSS"],
        ["LV用API・振り分け・会話制御", "独自実装", "自作（研究の中核）"],
    ], Inches(0.8), Inches(1.35), Inches(11.7), col_widths=[4, 6, 3],
        row_h=Inches(0.72), head_h=Inches(0.5), size=16, head_size=17, highlight_row=5)
    _, tf = textbox(s, Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.8))
    para(tf, "Dify / Flowise / LiteLLM / NextChat / OpenWebUI などの既製LLM製品は不使用",
         size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True)

    # 補4
    s = base_slide(prs, "補4．発言順の制御")
    _, tf = textbox(s, Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    para(tf, "「誰が次に話すかを決める処理」と「実際に呼び出す処理」を分離してある",
         size=19, color=GRAY, first=True)
    card(s, Inches(0.8), Inches(1.9), Inches(5.6), Inches(2.3), "現在（段階B）", [
        "順番を固定の規則で決める（巡回方式）。",
        "決めた順に1体ずつ呼び出し、",
        "直前までの発言を文脈として渡す。"], accent=ACCENT)
    card(s, Inches(7.0), Inches(1.9), Inches(5.6), Inches(2.3), "将来（段階C）", [
        "司会AIが文脈を読んで次の発話者を決める。",
        "「決める処理」だけを差し替えればよく、",
        "呼び出し側の実装は変更不要。"], accent=RED)
    _, tf = textbox(s, Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.8))
    para(tf, "部分成功の扱い", size=20, bold=True, color=NAVY, first=True, space_after=8)
    para(tf, "1体が失敗しても中断せず、残りのAIの発言は返す。"
             "どの発言が失敗したかは個別に判別できるようにしている。", size=17)

    # 補5
    s = base_slide(prs, "補5．堅牢性への配慮")
    table(s, [
        ["項目", "対策", "理由"],
        ["入力の長さ", "8,000文字で打ち切り", "過大な入力によるコスト・遅延の増加を防ぐ"],
        ["会話履歴", "直近20件までに制限", "履歴の肥大による遅延を防ぐ"],
        ["AIの数・巡回数", "上限を設定（最大6発言）", "応答時間が際限なく伸びるのを防ぐ"],
        ["応答待ち", "60秒でタイムアウト", "無応答時に処理が止まり続けるのを防ぐ"],
        ["1体の失敗", "中断せず残りを継続", "1体の障害で全体が落ちないようにする"],
        ["自動テスト", "LLMを模擬して検証", "APIを呼ばずに動作を確認できる"],
    ], Inches(0.8), Inches(1.35), Inches(11.7), col_widths=[3, 4, 6],
        row_h=Inches(0.72), head_h=Inches(0.5), size=15, head_size=16)

    # 補6
    s = base_slide(prs, "補6．開発体制")
    card(s, Inches(0.8), Inches(1.35), Inches(3.7), Inches(2.3), "ブランチ運用", [
        "main：動作確認済みの版",
        "develop：統合用",
        "feature：作業単位",
        "",
        "壊れても戻せる状態を保つ"], head_size=19, body_size=15)
    card(s, Inches(4.8), Inches(1.35), Inches(3.7), Inches(2.3), "ドキュメント", [
        "仕様書・構成図",
        "LabVIEW構築手順書",
        "研究室向け導入手順書",
        "設計判断の記録（ADR）"], head_size=19, body_size=15)
    card(s, Inches(8.8), Inches(1.35), Inches(3.7), Inches(2.3), "検証", [
        "自動テスト（LLMは模擬）",
        "疎通確認スクリプト",
        "実機での動作確認"], head_size=19, body_size=15)
    _, tf = textbox(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.0))
    para(tf, "引き継ぎを前提とした整備", size=20, bold=True, color=NAVY, first=True, space_after=10)
    para(tf, "設計の判断理由をADRとして記録しており、なぜその構成にしたかを後から追える。"
             "導入手順はDockerを前提に整理してあり、Node.jsやGitが無いPCでも起動できる。", size=17)

    os.makedirs(EXPORT, exist_ok=True)
    prs.save(OUT)
    print("生成しました:", OUT)
    print("スライド枚数:", len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == "__main__":
    build()
