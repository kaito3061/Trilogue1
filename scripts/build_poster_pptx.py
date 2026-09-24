#!/usr/bin/env python3
"""ポスター用（A4横に印刷して貼る）8枚版を生成する。

先生の指示:
    「タイトルと目次で1枚、ほかで7枚の合計8枚に調整しましょう」
    「A1に出す必要はない。A4横のまま印刷して貼れば解決する」

発表用の20枚（build_midterm_pptx.py）から、貼って読ませる前提で内容を絞る。
動画は印刷できないため、動作結果は静止画に差し替える。

    python3 scripts/build_poster_pptx.py
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_midterm_pptx as bm
from build_midterm_pptx import (
    ACCENT,
    EXPORT,
    GRAY,
    LIGHT,
    NAVY,
    RED,
    SH,
    SW,
    WHITE,
    RGBColor,
    base_slide,
    card,
    para,
    picture,
    roadmap,
    table,
    textbox,
)

OUT = os.path.join(EXPORT, "中間発表_柴尾_ポスター8枚.pptx")

# 動作結果に貼る静止画。動画から切り出したもの（3体の返答が映っている場面）。
bm.IMG["demo"] = os.path.join(EXPORT, "demo_multi_still.png")


def sheet(prs, title, number):
    """ポスターは1枚ずつ独立して読まれるため、枚数を「1/8」の形で入れる。"""
    s = base_slide(prs, title)
    _, tf = textbox(s, SW - Inches(1.3), SH - Inches(0.55), Inches(1.0), Inches(0.4))
    para(tf, f"{number} / 8", size=13, color=GRAY, align=PP_ALIGN.RIGHT, first=True)
    return s


def lead(slide, text, y=Inches(1.05), size=17, color=GRAY):
    """見出し直下の一行。何のスライドかを単独で分かるようにする。"""
    _, tf = textbox(slide, Inches(0.7), y, Inches(11.9), Inches(0.5))
    para(tf, text, size=size, color=color, first=True)


def band(slide, x, y, w, h, text, size=22, fill=NAVY, color=WHITE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, first=True)
    return box


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # ---------- 1/8 表紙＋目次 ----------
    s = base_slide(prs)
    top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(3.3))
    top.fill.solid()
    top.fill.fore_color.rgb = NAVY
    top.line.fill.background()
    top.shadow.inherit = False
    _, tf = textbox(s, Inches(0.7), Inches(0.5), Inches(11.9), Inches(2.4))
    para(tf, "既存計測系と大規模言語モデルを繋ぐゲートウェイの構築", size=34, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, first=True, space_after=10)
    para(tf, "― LabVIEWによる複数AI対話の実現 ―", size=23,
         color=RGBColor(0xC8, 0xDA, 0xE6), align=PP_ALIGN.CENTER, space_after=18)
    para(tf, "情報システム工学科　松田研究室　　2022531028　柴尾 海翔", size=20,
         color=WHITE, align=PP_ALIGN.CENTER, space_after=6)
    para(tf, "2026年 9月 28日　卒業研究 中間発表", size=15,
         color=RGBColor(0xC8, 0xDA, 0xE6), align=PP_ALIGN.CENTER)

    _, tf = textbox(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.5))
    para(tf, "本ポスターの構成", size=20, bold=True, color=NAVY, first=True)

    toc = [
        ("2", "研究背景と問題点", "主観は測れていない／画面変更のたびに作り直し"),
        ("3", "研究目的", "LLMを利用できる接続基盤と、対話構造の比較"),
        ("4", "システム構成と設計方針", "LabVIEW側の実装を増やさない3つの工夫"),
        ("5", "マルチエージェント化の段階", "段階A / B / C と現在地"),
        ("6", "動作結果", "1回の送信で3体のAIが順番に発言する"),
        ("7", "検証計画", "条件A〜Dを比較し、適した対話構造を明らかにする"),
        ("8", "今後の方針と本研究の位置づけ", "既存研究とLLMを繋ぐゲートウェイ"),
    ]
    y = Inches(4.15)
    for num, head, body in toc:
        n_box = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), y, Inches(0.42),
                                   Inches(0.42))
        n_box.fill.solid()
        n_box.fill.fore_color.rgb = ACCENT
        n_box.line.fill.background()
        n_box.shadow.inherit = False
        tfn = n_box.text_frame
        tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tfn, num, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)

        _, tf = textbox(s, Inches(1.45), y - Inches(0.03), Inches(11.0), Inches(0.5))
        p = para(tf, head + "　", size=19, bold=True, color=NAVY, first=True)
        run = p.add_run()
        run.text = body
        bm.set_font(run, size=15, color=GRAY)
        y += Inches(0.44)

    # ---------- 2/8 研究背景と問題点 ----------
    s = sheet(prs, "研究背景と問題点", 2)
    lead(s, "リハビリテーション評価は療法士の目視採点が主流で、対象者の増加により負担も増している。")
    card(s, Inches(0.7), Inches(1.6), Inches(5.9), Inches(1.9), "客観データ（測れている）",
         ["センサーで数値として取得できる",
          "指の屈曲角 ・ 運動回数 ・ 脈波"], accent=ACCENT, head_size=19, body_size=15)
    card(s, Inches(6.9), Inches(1.6), Inches(5.7), Inches(1.9), "主観（測れていない）",
         ["「動きやすくなった実感はあるか」",
          "「続けたいと思えるか」"], accent=RED, head_size=19, body_size=15)
    band(s, Inches(0.7), Inches(3.65), Inches(11.9), Inches(0.75),
         "両方が揃って、はじめて評価になる", size=22)
    card(s, Inches(0.7), Inches(4.65), Inches(5.9), Inches(1.8),
         "問題点 ①　主観を引き出す仕組みがない",
         ["・質問紙は形式が固定されており、その場の話に",
          "　応じた掘り下げができない",
          "・人が聞き取ると、聞き手の関心で内容が偏る"],
         accent=RED, head_size=18, body_size=14)
    card(s, Inches(6.9), Inches(4.65), Inches(5.7), Inches(1.8),
         "問題点 ②　画面変更のたびに作り直し",
         ["・計測画面はLabVIEWで作られており、",
          "　評価内容が変わるたびに画面も変わる",
          "・別にWeb画面を作ると二重の修正が必要になる"],
         accent=RED, head_size=18, body_size=14)
    _, tf = textbox(s, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.5))
    para(tf, "→ 偏りなく主観を引き出せて、かつ画面の変更に強い仕組みが要る",
         size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True)

    # ---------- 3/8 研究目的 ----------
    s = sheet(prs, "研究目的", 3)
    band(s, Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.1),
         "本研究：LabVIEWから大規模言語モデル（LLM）を利用できる接続基盤を構築する", size=21)
    band(s, Inches(0.8), Inches(2.45), Inches(11.7), Inches(1.1),
         "最終目的：主観情報を多面的かつ効率的に収集できるLLM対話構造を明らかにする",
         size=20, fill=LIGHT, color=ACCENT)
    for i, (head, body) in enumerate([
        ("① 画面変更に強い構成", ["LabVIEWをクライアントとし、",
                                   "画面の変更をLabVIEW側だけで",
                                   "完結できるようにする"]),
        ("② 対話構造を試せる基盤", ["複数のAIが順番に発言する",
                                     "会話をサーバー側で制御し、",
                                     "構成を差し替えられるようにする"]),
        ("③ 引き継げる環境", ["研究室の他のPCでも",
                               "同じように動く構成にする"]),
    ]):
        card(s, Inches(0.8) + i * Inches(4.0), Inches(3.75), Inches(3.7), Inches(2.2),
             head, body, head_size=19, body_size=15)
    _, tf = textbox(s, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.5))
    para(tf, "AIを何体どう並べるのが良いかは自明ではないため、構成を比較して見極める必要がある",
         size=16, color=GRAY, align=PP_ALIGN.CENTER, first=True)

    # ---------- 4/8 システム構成と設計方針 ----------
    s = sheet(prs, "システム構成と設計方針", 4)
    picture(s, "asis", Inches(0.5), Inches(1.15), Inches(12.3), Inches(1.75),
            caption="現在のシステム構成（As-Is）")
    _, tf = textbox(s, Inches(0.7), Inches(3.15), Inches(11.9), Inches(0.45))
    para(tf, "LabVIEW側の実装を増やさないための3つの工夫", size=19, bold=True,
         color=NAVY, first=True)
    table(s, [
        ["工夫", "内容", "LabVIEW側での効果"],
        ["① 平坦なJSON", "入れ子構造をやめ、フラットな形にする",
         "クラスタを1つ作るだけで送受信できる"],
        ["② 常にHTTP 200", "成否は本文の ok フラグで表す",
         "ステータス分岐が不要。Case Structure 1つで済む"],
        ["③ 会話を1本の文字列でも返す", "発言の配列に加え、連結済みの文字列も返す",
         "配列を解析せずに会話全体を表示できる"],
    ], Inches(0.7), Inches(3.7), Inches(11.9), col_widths=[3, 5, 5.5],
        row_h=Inches(0.82), head_h=Inches(0.45), size=15, head_size=16)
    _, tf = textbox(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5))
    para(tf, "→ 複雑さをサーバー側に寄せることで、画面の変更にLabVIEW側だけで追従できる",
         size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True)

    # ---------- 5/8 マルチエージェント化の段階 ----------
    s = sheet(prs, "マルチエージェント化の段階", 5)
    picture(s, "tobe", Inches(0.6), Inches(1.15), Inches(6.5), Inches(5.4),
            caption="将来構成（To-Be）")
    table(s, [
        ["段階", "内容", "状態"],
        ["A", "LabVIEW主導でAIを1体ずつ呼ぶ", "実現可能"],
        ["B", "サーバーが発言順を制御しまとめて返す", "実装完了 ←現在地"],
        ["C", "司会AIが発言順を動的に決める", "次の課題"],
    ], Inches(7.5), Inches(1.4), Inches(5.1), col_widths=[1, 5, 2.4],
        row_h=Inches(0.85), head_h=Inches(0.45), size=14, head_size=15, highlight_row=2)
    card(s, Inches(7.5), Inches(4.3), Inches(5.1), Inches(1.8), "段階Cへの備え",
         ["発言順を「決める処理」と「実行する処理」を",
          "分けて実装した。司会AIを入れるときは、",
          "決める部分だけを差し替えれば済む。"],
         head_size=18, body_size=14)

    # ---------- 6/8 動作結果 ----------
    s = sheet(prs, "動作結果　LabVIEWから1回送るだけで3体が会話する", 6)
    picture(s, "demo", Inches(0.6), Inches(1.15), Inches(6.6), Inches(3.9),
            caption="LabVIEWの画面に返ってきた3体の返答（動画から抜粋）")
    card(s, Inches(0.6), Inches(5.3), Inches(6.6), Inches(1.5), "入力（1文だけ）",
         ["「リハビリのトレーニングを終えました。",
          "　指が少し動きやすくなった気がします。」"],
         head_size=17, body_size=15)
    table(s, [
        ["順", "AI", "発言の要旨", "所要"],
        ["1", "Alpha", "改善が見られるのは励みになる。特に効果を感じた"
                       "トレーニングはあったかと聞き返す", "2.0秒"],
        ["2", "Beta", "継続的な取り組みと適切なフィードバックが必要。"
                      "共有し合える環境を整えることも大切", "2.9秒"],
        ["3", "Gamma", "具体的な目標を設定すると効果を引き出せる。"
                       "日常生活で指を使う作業への挑戦を提案", "1.1秒"],
    ], Inches(7.5), Inches(1.3), Inches(5.1), col_widths=[0.6, 1.3, 5.2, 1.1],
        row_h=Inches(1.15), head_h=Inches(0.42), size=12, head_size=13, highlight_row=2)
    card(s, Inches(7.5), Inches(5.0), Inches(5.1), Inches(1.8), "注目点", [
        "1体目は問い返し、2体目は条件を補い、",
        "3体目は次の行動へ広げている。",
        "前の発言を読んだうえで話すため重複しない。",
        "※ この3体構成は比較条件の一つ（条件B）。"],
        accent=RED, head_size=18, body_size=14)

    # ---------- 7/8 検証計画 ----------
    s = sheet(prs, "検証計画 ― 主観情報の収集に適した対話構造の比較", 7)
    lead(s, "3体が最適という主張ではない。現在の構成は比較条件の一つであり、A〜Dを比較して"
            "適した構造を明らかにする。")
    table(s, [
        ["条件", "LLM構成", "狙い", "本研究での状態"],
        ["A", "LLM 1体", "基準条件", "実装済み（単体AI版）"],
        ["B", "1体＋複数ペルソナ", "視点を変える効果", "実装済み ← 今回の動作結果"],
        ["C", "複数LLM＋固定順", "独立した役割を持つ効果", "設定変更で対応可能"],
        ["D", "複数LLM＋司会AI", "対話を動的に制御する効果", "次段階で実装"],
    ], Inches(0.7), Inches(1.7), Inches(11.9), col_widths=[1, 3.2, 3.6, 4],
        row_h=Inches(0.72), head_h=Inches(0.5), size=16, head_size=17, highlight_row=2)
    card(s, Inches(0.7), Inches(5.0), Inches(5.8), Inches(1.6), "応答時間（AI 3体の実測）",
         ["約 6.0 〜 6.3 秒（2.0 / 2.9 / 1.1 秒）",
          "うち自作サーバーの処理は 7 ミリ秒。",
          "遅延のほぼ全てがLLM側の応答待ち。"],
         accent=ACCENT, head_size=18, body_size=14)
    card(s, Inches(6.8), Inches(5.0), Inches(5.8), Inches(1.6), "比較で見たいこと",
         ["・重複のない多面的な意見が得られるか",
          "・応答時間と発言数のつり合い",
          "・患者の負担にならない対話の長さ"],
         accent=RED, head_size=18, body_size=14)

    # ---------- 8/8 今後の方針と本研究の位置づけ ----------
    s = sheet(prs, "今後の方針と本研究の位置づけ", 8)
    roadmap(s, Inches(0.6), Inches(1.15), Inches(12.1), Inches(2.0))
    cols = [
        (Inches(0.6), Inches(3.5), "松田研の既存研究", [
            "・リハビリ用デバイス（力覚・屈曲データ）",
            "・ハードウェア制御（PWM / AD・DA変換）",
            "・電動車椅子・計測系",
            "いずれもLabVIEWが窓口"], GRAY),
        (Inches(4.9), Inches(3.6), "本研究：接続基盤", [
            "・LabVIEWから1回投げるだけ",
            "・サーバーが会話を制御",
            "・LabVIEW側に追加実装は不要",
            "計測系を作り変えずにLLMを利用できる"], ACCENT),
        (Inches(9.3), Inches(3.4), "LLM / マルチエージェント", [
            "・複数AIによる対話",
            "・主観情報の引き出し",
            "・司会AIによる発言順の制御",
            "（外部API / 将来は院内）"], NAVY),
    ]
    for x, w, head, lines, accent in cols:
        card(s, x, Inches(3.45), w, Inches(2.1), head, lines,
             accent=accent, head_size=17, body_size=13)
    for x in (Inches(4.45), Inches(8.85)):
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, Inches(4.25), Inches(0.4),
                                Inches(0.45))
        ar.fill.solid()
        ar.fill.fore_color.rgb = RGBColor(0xC8, 0xD3, 0xDF)
        ar.line.fill.background()
        ar.shadow.inherit = False
    band(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(0.75),
         "研究室の様々な研究にLLMを挿せる土台を作ったことに価値がある", size=21)
    _, tf = textbox(s, Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.5))
    para(tf, "※ 本研究の完了範囲は接続基盤と司会AIの導入まで。"
             "センサーデータとの統合以降は発展の方向性として示している。",
         size=14, color=RED, align=PP_ALIGN.CENTER, first=True)

    prs.save(OUT)
    print(f"生成しました: {OUT}")
    print(f"スライド枚数: {len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    build()
